from __future__ import annotations

from dataclasses import dataclass, field
import re
from typing import Iterable

import docx
import pdfplumber
from pypdf import PdfReader


_HEADING_NUMBER_RE = re.compile(r"^\s*\d+(?:\.\d+)*\s+(.+)$")


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int | None
    text: str
    tables: tuple[str, ...] = ()
    headings: tuple[str, ...] = ()
    extraction_method: str = "text"
    has_images: bool = False
    ocr_required: bool = False

    @property
    def content(self) -> str:
        return "\n\n".join(
            part.strip()
            for part in (self.text, *self.tables)
            if part and part.strip()
        )


@dataclass(frozen=True)
class ExtractedDocument:
    pages: tuple[ExtractedPage, ...]
    file_type: str
    extraction_method: str
    metadata: dict[str, object] = field(default_factory=dict)

    @property
    def content(self) -> str:
        return "\n\n".join(page.content for page in self.pages if page.content)


def _table_markdown(rows: Iterable[Iterable[object | None]]) -> str:
    clean_rows: list[list[str]] = []
    for row in rows:
        cells = [
            str(cell or "").replace("\n", " ").replace("|", "\\|").strip()
            for cell in row
        ]
        if any(cells):
            clean_rows.append(cells)
    if not clean_rows:
        return ""

    width = max(len(row) for row in clean_rows)
    clean_rows = [row + ([""] * (width - len(row))) for row in clean_rows]
    lines = ["| " + " | ".join(clean_rows[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")
    lines.extend("| " + " | ".join(row) + " |" for row in clean_rows[1:])
    return "\n".join(lines)


def _inside_table_bbox(obj: dict, bboxes: list[tuple[float, float, float, float]]) -> bool:
    try:
        midpoint_x = (float(obj["x0"]) + float(obj["x1"])) / 2
        midpoint_y = (float(obj["top"]) + float(obj["bottom"])) / 2
    except (KeyError, TypeError, ValueError):
        return False
    return any(
        x0 <= midpoint_x <= x1 and top <= midpoint_y <= bottom
        for x0, top, x1, bottom in bboxes
    )


def _detected_headings(text: str) -> tuple[str, ...]:
    headings: list[str] = []
    for raw_line in str(text or "").splitlines():
        line = raw_line.strip().lstrip("#").strip()
        if not line or len(line) > 120:
            continue
        letters = [character for character in line if character.isalpha()]
        uppercase_ratio = (
            sum(character.isupper() for character in letters) / len(letters)
            if letters
            else 0.0
        )
        if raw_line.lstrip().startswith("#") or _looks_numbered_heading(line) or (
            len(letters) >= 4 and uppercase_ratio >= 0.8
        ):
            if line not in headings:
                headings.append(line)
    return tuple(headings)


def _looks_numbered_heading(text: str) -> bool:
    match = _HEADING_NUMBER_RE.match(text)
    if not match:
        return False
    title = match.group(1).strip()
    words = re.findall(r"[A-Za-z]+", title)
    if not words or len(words) > 8 or len(title) > 80:
        return False
    significant = [word for word in words if word.lower() not in {"a", "an", "and", "for", "of", "the", "to"}]
    return bool(
        significant
        and sum(word[0].isupper() for word in significant) / len(significant) >= 0.75
    )


def _ocr_required(text: str, has_images: bool) -> bool:
    alphanumeric = sum(character.isalnum() for character in str(text or ""))
    return bool(has_images and alphanumeric < 40)


def extract_pdf(file_path: str) -> ExtractedDocument:
    pages: list[ExtractedPage] = []
    with pdfplumber.open(file_path) as plumber_pdf:
        for page_number, page in enumerate(plumber_pdf.pages, start=1):
            table_markdown: list[str] = []
            table_bboxes: list[tuple[float, float, float, float]] = []
            try:
                tables = page.find_tables().tables
            except Exception:
                tables = []

            for table in tables:
                rendered = _table_markdown(table.extract() or [])
                if rendered:
                    table_markdown.append(rendered)
                    table_bboxes.append(tuple(float(value) for value in table.bbox))

            text_page = page
            if table_bboxes:
                text_page = page.filter(
                    lambda obj: not _inside_table_bbox(obj, table_bboxes)
                )
            plain_text = text_page.extract_text() or ""
            has_images = bool(page.images)
            combined = "\n\n".join([plain_text, *table_markdown])
            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    text=plain_text.strip(),
                    tables=tuple(table_markdown),
                    headings=_detected_headings(plain_text),
                    extraction_method="pdfplumber",
                    has_images=has_images,
                    ocr_required=_ocr_required(combined, has_images),
                )
            )
    return ExtractedDocument(
        pages=tuple(pages),
        file_type="pdf",
        extraction_method="pdfplumber",
    )


def extract_pdf_fast(file_path: str) -> ExtractedDocument:
    reader = PdfReader(file_path)
    pages: list[ExtractedPage] = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        try:
            has_images = bool(page.images)
        except Exception:
            has_images = False
        pages.append(
            ExtractedPage(
                page_number=page_number,
                text=page_text.strip(),
                headings=_detected_headings(page_text),
                extraction_method="pypdf",
                has_images=has_images,
                ocr_required=_ocr_required(page_text, has_images),
            )
        )
    return ExtractedDocument(
        pages=tuple(pages),
        file_type="pdf",
        extraction_method="pypdf",
    )


def extract_docx(file_path: str) -> ExtractedDocument:
    document = docx.Document(file_path)
    parts: list[str] = []
    headings: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        parts.append(text)
        style_name = str(getattr(paragraph.style, "name", "") or "").lower()
        if style_name.startswith("heading"):
            headings.append(text)

    tables: list[str] = []
    for table in document.tables:
        rendered = _table_markdown(
            [[cell.text for cell in row.cells] for row in table.rows]
        )
        if rendered:
            tables.append(rendered)

    text = "\n\n".join(parts)
    return ExtractedDocument(
        pages=(
            ExtractedPage(
                page_number=None,
                text=text,
                tables=tuple(tables),
                headings=tuple(dict.fromkeys([*headings, *_detected_headings(text)])),
                extraction_method="python-docx",
            ),
        ),
        file_type="docx",
        extraction_method="python-docx",
    )


def extract_text_file(file_path: str, file_type: str) -> ExtractedDocument:
    with open(file_path, "r", encoding="utf-8", errors="replace") as file:
        text = file.read()
    return extracted_document_from_text(text, file_type=file_type)


def extracted_document_from_text(
    text: str,
    *,
    file_type: str = "txt",
    extraction_method: str = "text",
) -> ExtractedDocument:
    return ExtractedDocument(
        pages=(
            ExtractedPage(
                page_number=None,
                text=str(text or ""),
                headings=_detected_headings(text),
                extraction_method=extraction_method,
            ),
        ),
        file_type=file_type,
        extraction_method=extraction_method,
    )


def load_pdf(file_path: str) -> str:
    return extract_pdf(file_path).content


def load_pdf_fast(file_path: str) -> str:
    """Fast raw text extraction using pypdf."""
    return extract_pdf_fast(file_path).content


def load_docx(file_path: str) -> str:
    return extract_docx(file_path).content


def load_txt(file_path: str) -> str:
    return extract_text_file(file_path, "txt").content


def extract_document(
    file_path: str,
    file_type: str,
    strategy: str = "auto",
) -> ExtractedDocument:
    normalized_type = str(file_type or "").lower().lstrip(".")
    if normalized_type == "pdf":
        return extract_pdf_fast(file_path) if strategy == "fast" else extract_pdf(file_path)
    if normalized_type == "docx":
        return extract_docx(file_path)
    if normalized_type in {"txt", "md"}:
        return extract_text_file(file_path, normalized_type)
    raise ValueError("Unsupported file type")


def load_document(file_path: str, file_type: str, strategy: str = "auto") -> str:
    return extract_document(file_path, file_type, strategy).content
